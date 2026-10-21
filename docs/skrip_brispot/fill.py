import sys, difflib, re
from load import *   # uses sys.argv[1] as source
import openpyxl
SRC=sys.argv[1]; OUT=sys.argv[2]
m=re.search(r'(\d{1,2})[ _]([A-Za-z]{3})[ _](\d{2})[ _](\d{1,2})[.:](\d{2})',SRC)
jam=f"{m.group(4).zfill(2)}:{m.group(5)}" if m else ''
nom,rm,mt,ren_rm,ren_mt,tgl=load()
tgl=pd.Timestamp(tgl).normalize()
BLN=['Januari','Februari','Maret','April','Mei','Juni','Juli','Agustus','September','Oktober','November','Desember']
tgl_id=f"{tgl.day:02d} {BLN[tgl.month-1]} {tgl.year}"
rmk=set(rm.NAMA.map(norm)); mk=set(mt.NAMA.map(norm)); master=list(rmk|mk)
alias={n:difflib.get_close_matches(n,master,1,0.9)[0] for n in nom.KEY_NAMA.unique() if n not in master and difflib.get_close_matches(n,master,1,0.9)}
nom['KEY_NAMA']=nom.KEY_NAMA.replace(alias)
u=nom[nom.DOUBLE==0]
def stats(k):
    g=u[u.KEY_NAMA==k]
    def ds(mask,col='PLAFON_N'): x=g[mask]; return len(x), x[col].sum()/1e6
    w=ds(g.STATUS=='PRAKARSA'); p=ds(g.STATUS=='MENUNGGU PUTUSAN'); a=ds(g.STATUS=='APPROVED'); c=ds(g.STATUS=='MENUNGGU PENCAIRAN')
    cs=g[(g.STATUS=='PENCAIRAN SUKSES')&(g.TGL_CAIR==tgl)]
    return dict(wd=w[0],wp=w[1],pd=p[0],pp=p[1],ad=a[0],ap=a[1],md=c[0],mp=c[1],sd=len(cs),sp=cs.PLAFON_N.sum()/1e6,sn=cs['NET OS_N'].sum()/1e6)

wb=openpyxl.load_workbook(SRC); wv=openpyxl.load_workbook(SRC,data_only=True)
KEEP=['RM BGRITEL','KC BGRITEL','Mantri BGMIKRO','Unit BGMIKRO']
for n in list(wb.sheetnames):
    if n not in KEEP: del wb[n]
for n in list(wb.defined_names): del wb.defined_names[n]
def total_rows(ws):
    return [r for r in range(1,ws.max_row+1) if str(ws.cell(r,1).value).strip().upper()=='TOTAL']
def wa(seg,lvl): return f"Berikut disampaikan Brispot Briguna {seg} per {lvl} posisi {tgl_id} Brispot (Update {jam} WIB)."
def fix_static(ws, cells):
    for c,v in cells.items(): ws[c].value=v

# ---------- per orang ----------
def person_sheet(name, off, renmap, keycol_idx):
    # off: column offset (0 for RM layout starting E=Rencana; 2 for Mantri layout starting G)
    ws=wb[name]; vs=wv[name]
    t1,t2=total_rows(ws)
    top=range(7,t1); hdr2=[r for r in range(t1+1,t2) if ws.cell(r,1).value=='No'][0]; bot=range(hdr2+3,t2)
    C=lambda i: i+off   # col index helper: RM E=5
    rows=[]
    for i,r in enumerate(top,1):
        nm=vs.cell(r,keycol_idx).value; k=norm(nm)
        s=stats(k); rd,rn=renmap.get(k,(0,0))
        info=[vs.cell(r,j).value for j in range(2,keycol_idx+1+(1 if off else 0))]  # B..name (+D/H for mantri)
        q=vs.cell(r,17+off).value or 0
        rows.append(dict(info=info,rd=rd,rn=rn,q=q,**s))
        ws.cell(r,1).value=i
        for j,v in enumerate(info,2): ws.cell(r,j).value=v
        vals={C(5):rd,C(6):rn,C(7):s['wd'],C(8):s['wp'],C(9):s['pd'],C(10):s['pp'],C(11):s['ad'],C(12):s['ap'],C(13):s['md'],C(14):s['mp'],C(19):s['sd'],C(20):s['sp'],C(21):s['sn']}
        # template columns: E rd F rn G wd H wp I pd J pp K ad L ap M md N mp O/P total Q guid R gap S sd T sp U sn V gap
        vals={5+off:rd,6+off:rn,7+off:s['wd'],8+off:s['wp'],9+off:s['pd'],10+off:s['pp'],11+off:s['ad'],12+off:s['ap'],13+off:s['md'],14+off:s['mp'],19+off:s['sd'],20+off:s['sp'],21+off:s['sn']}
        for c,v in vals.items(): ws.cell(r,c).value=v
        ws.cell(r,15+off).value=f"={openpyxl.utils.get_column_letter(13+off)}{r}+{openpyxl.utils.get_column_letter(11+off)}{r}+{openpyxl.utils.get_column_letter(9+off)}{r}"
        ws.cell(r,16+off).value=f"={openpyxl.utils.get_column_letter(14+off)}{r}+{openpyxl.utils.get_column_letter(12+off)}{r}+{openpyxl.utils.get_column_letter(10+off)}{r}"
    # bottom (values, sorted)
    for x in rows: x['hp']=x['pp']+x['ap']+x['mp']; x['hd']=x['pd']+x['ad']+x['md']
    srt=sorted(rows,key=lambda x:(-x['sn'],-x['hp'],-x['wp']))
    for i,(r,x) in enumerate(zip(bot,srt),1):
        ws.cell(r,1).value=i
        for j,v in enumerate(x['info'],2): ws.cell(r,j).value=v
        vals=[x['rd'],x['rn'],x['wd'],x['wp'],x['pd'],x['pp'],x['ad'],x['ap'],x['md'],x['mp'],x['hd'],x['hp'],x['q'],x['hp']-x['q'],x['sd'],x['sp'],x['sn'],x['sn']-x['rn']]
        for j,v in enumerate(vals,5+off): ws.cell(r,j).value=v
    return rows,hdr2

R,h_rm=person_sheet('RM BGRITEL',0,ren_rm,4)
M,h_mt=person_sheet('Mantri BGMIKRO',2,ren_mt,5)

# ---------- per uker (top tetap rumus ke sheet orang; bottom nilai) ----------
def uker_sheet(name, rows, keyfn, off, ncount_col, lookup_static, info_fn):
    ws=wb[name]; vs=wv[name]
    t1,t2=total_rows(ws)
    top=range(7,t1); hdr2=[r for r in range(t1+1,t2) if ws.cell(r,1).value=='No'][0]; bot=range(hdr2+3,t2)
    agg={}
    for x in rows:
        k=keyfn(x); a=agg.setdefault(k,{'n':0,**{f:0 for f in ['rd','rn','wd','wp','pd','pp','ad','ap','md','mp','hd','hp','q','sd','sp','sn']}})
        a['n']+=1
        for f in a:
            if f!='n': a[f]+=x[f]
    keys=[]
    for i,r in enumerate(top,1):
        ws.cell(r,1).value=i
        lookup_static(ws,vs,r)
        keys.append(info_fn(vs,r))
    srt=sorted(keys,key=lambda k:(-agg[k[-1]]['sn'],-agg[k[-1]]['hp'],-agg[k[-1]]['wp']))
    for i,(r,k) in enumerate(zip(bot,srt),1):
        a=agg[k[-1]]; ws.cell(r,1).value=i
        for j,v in enumerate(k,3): ws.cell(r,j).value=v
        vals=[a['n'],a['rd'],a['rn'],a['wd'],a['wp'],a['pd'],a['pp'],a['ad'],a['ap'],a['md'],a['mp'],a['hd'],a['hp'],a['q'],a['hp']-a['q'],a['sd'],a['sp'],a['sn'],a['sn']-a['rn']]
        for j,v in enumerate(vals,ncount_col): ws.cell(r,j).value=v
    return hdr2
def kc_static(ws,vs,r): pass
h_kc=uker_sheet('KC BGRITEL',R,lambda x:x['info'][1],0,4,kc_static,lambda vs,r:(vs.cell(r,3).value,))
def unit_static(ws,vs,r):
    ws.cell(r,3).value=vs.cell(r,3).value; ws.cell(r,4).value=vs.cell(r,4).value
h_un=uker_sheet('Unit BGMIKRO',M,lambda x:x['info'][2],0,5,unit_static,lambda vs,r:(vs.cell(r,3).value,vs.cell(r,4).value))

# ---------- teks statis (tanggal/jam tarikan) ----------
posisi=f"Posisi tanggal :{tgl:%d %b %Y}"
pukul=f"Pukul : {jam}"
fix_static(wb['RM BGRITEL'],{'X1':tgl.to_pydatetime(),'Y1':jam,'A2':posisi,'A52':posisi,'D52':pukul,'X53':wa('Ritel','RM')})
fix_static(wb['KC BGRITEL'],{'X1':tgl.to_pydatetime(),'A3':posisi,'A35':posisi,'D35':pukul,'X36':wa('Ritel','KC')})
fix_static(wb['Mantri BGMIKRO'],{'Z1':tgl.to_pydatetime(),'AA1':jam,'A2':posisi,'A47':posisi,'E47':pukul,'Z48':wa('Mikro','Mantri')})
fix_static(wb['Unit BGMIKRO'],{'Y1':tgl.to_pydatetime(),'Z1':jam,'A2':posisi,'A42':posisi,'D42':pukul,'Y43':wa('Mikro','Unit')})
# kolom cek Mantri (Z7:Z38 '=H7=AL7') -> hapus
ws=wb['Mantri BGMIKRO']
for r in range(7,39):
    if isinstance(ws.cell(r,26).value,str) and ws.cell(r,26).value.startswith('=H'): ws.cell(r,26).value=None
# sisa rumus yg merujuk sheet/nama terhapus -> laporkan
bad=[]
for w in wb:
    for row in w.iter_rows():
        for c in row:
            v=getattr(c.value,'text',c.value)
            if isinstance(v,str) and v.startswith('=') and re.search(r'Nominatif!|LIST!|\bRM,|KANCA|NAMA_MANTRI|UNIT_KERJA|KC_UNIT|TODAY|NOW|SEQUENCE|FILTER|UNIQUE',v): bad.append((w.title,c.coordinate,v[:50]))
print('sisa rumus bermasalah:',bad[:10],len(bad))
wb.active=0
wb.save(OUT); print('alias',alias,'jam',jam,'tgl',tgl_id)
